#!/usr/bin/env python3
"""Build the content-map workbook from run.json.

usage: assemble.py run.json --profile client|internal --out file.xlsx

Layout mirrors the C+P Consulting reference exactly:
  IMPLEMENTATION PLAN | P1..Pn pillar tabs | All Pages (master)   (+ Keyword Research tab in the internal profile)
Language blocks are generated from client.markets, so the file works for one language or many.
"""
import json, math, sys, argparse
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
from page_types import Catalogue

HEAD_FILL = PatternFill('solid', fgColor='1F3864')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=10)
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(italic=True, size=10, color='555555')
SECTION_FILL = PatternFill('solid', fgColor='D9E2F3')
WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CAT = None  # set in build(): Catalogue(run), so custom types declared in run.json count


def vol_num(v):
    return v if isinstance(v, (int, float)) else 0


def priority(p):
    if p.get('priority_score') is not None:
        return p['priority_score']
    c = p['core']; v = vol_num(c.get('vol'))
    vp = 0 if v <= 0 else min(40, 8 * math.log10(v + 1) * 1.6)
    kd = c.get('kd'); dp = 20 if kd is None else max(0, 40 - 0.4 * kd)
    it = (c.get('intent') or '').lower(); ip = 15 if it in ('commercial', 'transactional') else 8 if it == 'informational' else 0
    tp = CAT.points(p.get('page_type'))
    score = vp + dp + ip + tp
    if v < 10 and not CAT.low_volume_ok(p.get('page_type')):
        score -= 10
    return int(max(0, min(100, round(score))))


def header(ws, row, cols):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row, i, name); c.font = HEAD_FONT; c.fill = HEAD_FILL; c.alignment = WRAP; c.border = BORDER
    ws.freeze_panes = ws.cell(row + 1, 1)


def widths(ws, spec):
    for i, w in enumerate(spec, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(run, profile, out):
    global CAT
    CAT = Catalogue(run)
    client = run['client']; markets = client['markets']
    primary = next(m for m in markets if m.get('primary')); others = [m for m in markets if not m.get('primary')]
    pcode = primary['code']
    pages = run['pages']; pillars = sorted(run['pillars'], key=lambda x: x['num'])
    internal = profile == 'internal'
    for p in pages:
        p['page_type'] = CAT.resolve(p.get('page_type'))   # old labels such as 'PILLAR HUB' still work
        p['priority_score'] = priority(p)

    # inbound links from links_out
    inbound = defaultdict(int); urls = {p['url'] for p in pages}
    for p in pages:
        for l in p.get('links_out', []):
            if l in urls and l != p['url']: inbound[l] += 1

    wb = Workbook(); wb.remove(wb.active)

    # ---- IMPLEMENTATION PLAN
    ws = wb.create_sheet('IMPLEMENTATION PLAN')
    ws.cell(1, 1, f"{client['name']}  |  Content Map and Implementation Plan").font = TITLE_FONT
    r = 2
    for sec in run.get('plan', {}).get('sections', []):
        c = ws.cell(r, 1, sec['title']); c.font = Font(bold=True, size=11); c.fill = SECTION_FILL
        ws.cell(r, 2).fill = SECTION_FILL; r += 1
        for label, text in sec['rows']:
            ws.cell(r, 1, label).font = Font(bold=True); ws.cell(r, 1).alignment = WRAP
            ws.cell(r, 2, text).alignment = WRAP; r += 1
        r += 1
    widths(ws, [28, 120])

    # ---- language columns
    def lang_cols(kind):
        cols = []
        for m in others:
            cols += [f"{m['language_name']} keyword", f"{m['code']} vol"]
            if internal: cols += [f"{m['code']} KD", f"{m['code']} YoY"]
        return cols

    def lang_vals(p):
        vals = []
        for m in others:
            L = (p.get('loc') or {}).get(m['code'], {})
            vals += [L.get('keyword'), L.get('vol_display', L.get('vol'))]
            if internal: vals += [L.get('kd'), L.get('yoy')]
        return vals

    # ---- pillar tabs
    by_pillar = defaultdict(list)
    for p in pages: by_pillar[p['pillar_num']].append(p)
    for pl in pillars:
        name = f"P{pl['num']} {pl['name']}"[:31]
        ws = wb.create_sheet(name)
        ws.cell(1, 1, f"P{pl['num']}. {pl['name']}").font = TITLE_FONT
        ws.cell(2, 1, f"{pl.get('entity','')}   |   Goal: {pl.get('goal','')}".strip()).font = SUB_FONT
        cols = ['Cluster', 'Page type', 'Core keyword', f'{pcode} vol/mo']
        if internal: cols += ['KD', 'YoY']
        cols += ['Search intent', 'Intent conf'] + lang_cols('pillar') + ['Priority score', 'Secondary keywords', 'H1', 'Meta title',
                 'H2 outline (researched)', 'Target AI prompts', 'Internal links OUT', 'Words']
        if not pl.get('glossary'): cols.append('Primary CTA')
        if internal: cols += ['URL slug', 'Page Maker type', 'Page status', 'Notes']
        header(ws, 4, cols)
        rr = 5
        for p in sorted(by_pillar[pl['num']], key=lambda x: (x['cluster'], -x['priority_score'])):
            c = p['core']
            vals = [p['cluster'], CAT.label(p['page_type']), c['keyword'], c.get('vol_display', c.get('vol'))]
            if internal: vals += [c.get('kd'), c.get('yoy')]
            vals += [c.get('intent'), c.get('intent_conf')] + lang_vals(p) + [p['priority_score'], p.get('secondaries'), p.get('h1'), p.get('meta'),
                     p.get('h2_outline'), p.get('ai_prompts'), ';'.join(p.get('links_out', [])) + ';' if p.get('links_out') else '', p.get('words')]
            if not pl.get('glossary'): vals.append(p.get('cta'))
            if internal: vals += [p['url'], p['page_type'], p.get('status', 'new'), p.get('notes')]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(rr, i, v); cell.alignment = WRAP; cell.border = BORDER
            rr += 1
        widths(ws, [22, 16, 26, 10] + ([7, 7] if internal else []) + [13, 10] + [22, 8, 7, 7][: (4 if internal else 2)] * len(others) + [9, 40, 40, 36, 60, 50, 50, 8, 30, 36, 20, 14, 60])

    # ---- master
    ws = wb.create_sheet('All Pages (master)')
    ws.cell(1, 1, 'All pages (master list)').font = TITLE_FONT
    ws.cell(2, 1, 'Every page in one filterable sheet. Filter by Publishing pod for a week\'s work. Sort by Inbound links ascending to find under-linked pages. Intent and confidence are researched values, not estimates.').font = SUB_FONT
    cols = ['Pillar #', 'Pillar', 'Cluster', 'Wave', 'Publishing pod', 'Priority score', 'Page type', 'URL slug', 'Core keyword', f'{pcode} vol/mo']
    if internal: cols += ['YoY', 'KD']
    cols += ['Search intent', 'Intent conf'] + lang_cols('master') + ['Secondary keywords', 'Words', 'H1', 'Meta title', 'Schema', 'Inbound links', 'Internal links OUT']
    if internal: cols += ['Page Maker type', 'Page status', 'Localised URLs', 'Answer block', 'Notes']
    header(ws, 4, cols)
    pname = {pl['num']: pl['name'] for pl in pillars}
    rr = 5
    def podkey(p):
        try: return int(str(p.get('pod', '')).split()[-1])
        except Exception: return 999
    for p in sorted(pages, key=lambda x: (podkey(x), -x['priority_score'])):
        c = p['core']
        vals = [p['pillar_num'], pname[p['pillar_num']], p['cluster'], p.get('wave'), p.get('pod'), p['priority_score'], CAT.label(p['page_type']), p['url'], c['keyword'], c.get('vol_display', c.get('vol'))]
        if internal: vals += [c.get('yoy'), c.get('kd')]
        vals += [c.get('intent'), c.get('intent_conf')] + lang_vals(p) + [p.get('secondaries'), p.get('words'), p.get('h1'), p.get('meta'), p.get('schema'), inbound.get(p['url'], 0), ';'.join(p.get('links_out', [])) + ';' if p.get('links_out') else '']
        if internal:
            loc_urls = ';'.join(f"/{m['language_code']}{p['url']}" for m in others)
            vals += [p['page_type'], p.get('status', 'new'), loc_urls, p.get('answer_block'), p.get('notes')]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(rr, i, v); cell.alignment = WRAP; cell.border = BORDER
        rr += 1
    ws.cell(rr, 7, 'TOTAL').font = Font(bold=True)
    ws.cell(rr, 8, f'=COUNTA(H5:H{rr-1})').font = Font(bold=True)
    tcol = cols.index('Internal links OUT') + 1
    ws.cell(rr, tcol - 1, f'=SUM({get_column_letter(tcol-1)}5:{get_column_letter(tcol-1)}{rr-1})').font = Font(bold=True)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{rr-1}"
    widths(ws, [7, 20, 22, 16, 12, 9, 16, 40, 26, 10] + ([7, 7] if internal else []) + [13, 10] + [22, 8, 7, 7][: (4 if internal else 2)] * len(others) + [40, 8, 40, 36, 34, 9, 60, 20, 12, 40, 60, 60])

    # ---- keyword research (internal only)
    if internal and run.get('keyword_research'):
        ws = wb.create_sheet('Keyword Research')
        ws.cell(1, 1, 'Keyword research: method and findings').font = TITLE_FONT
        r = 3
        for label, text in run['keyword_research'].get('findings', []):
            ws.cell(r, 1, label).font = Font(bold=True); ws.cell(r, 2, text).alignment = WRAP; r += 1
        r += 1
        for row in run['keyword_research'].get('native_table', []):
            for i, v in enumerate(row, 1): ws.cell(r, i, v)
            r += 1
        widths(ws, [30, 100])

    wb.save(out)
    print(json.dumps({'file': out, 'profile': profile, 'pages': len(pages), 'languages': [m['code'] for m in markets],
                      'urls_total': len(pages) * len(markets), 'orphans': [p['url'] for p in pages if inbound.get(p['url'], 0) == 0]}))


if __name__ == '__main__':
    ap = argparse.ArgumentParser(); ap.add_argument('run'); ap.add_argument('--profile', default='client', choices=['client', 'internal']); ap.add_argument('--out', required=True)
    a = ap.parse_args(); build(json.load(open(a.run)), a.profile, a.out)
